# -*- coding: utf-8 -*-
"""
Verificador de cotejo B1: comprueba que cada fila de la hoja "Diferencias"
del archivo cotejo_B1_diferencias.xlsx tenga su texto "escrito" presente,
tal cual, en la transcripcion real del minuto indicado. Calcula el
porcentaje de diferencias por tramo a partir del CONTEO REAL de filas,
no de una cifra escrita a mano.

Uso:
    python 07_Datos/scripts/plan_mejora/verificar_cotejo_B1.py \
        07_Datos/datos_procesados/cotejo_B1_diferencias.xlsx
"""
import sys
import re
from pathlib import Path
import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[3]
TRANSCRIPCIONES = REPO_ROOT / "02_Evidencias" / "Transcripciones"

ARCHIVO_POR_ENTR = {
    "ENTR-01": "2026-05-23_ENTR-01_Transcripcion.md",
    "ENTR-02": "2026-05-23_ENTR-02_Transcripcion.md",
    "ENTR-03": "2026-05-23_ENTR-03_Transcripcion.md",
    "ENTR-04": "2026-07-28_ENTR-04_Transcripcion.md",
    "ENTR-05": "2026-07-28_ENTR-05_Transcripcion.md",
    "ENTR-06": "2026-07-28_ENTR-06_Transcripcion.md",
    "ENTR-07": "2026-07-28_ENTR-07_Transcripcion.md",
    "ENTR-08": "2026-07-28_ENTR-08_Transcripcion.md",
}


def normalizar(s):
    s = s.lower()
    s = re.sub(r"[^\wáéíóúñü ]", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def minuto_a_segundos(m):
    mm, ss = m.split(":")
    return int(mm) * 60 + int(ss)


def cargar_transcripcion(entr):
    path = TRANSCRIPCIONES / ARCHIVO_POR_ENTR[entr]
    texto = path.read_text(encoding="utf-8")
    bloques = []
    for m in re.finditer(r"\[(\d{2}:\d{2})\][^\n]*?\*\*\s*([^\n]+)", texto):
        bloques.append((minuto_a_segundos(m.group(1)), m.group(2)))
    return bloques


def main():
    if len(sys.argv) < 2:
        print("Uso: verificar_cotejo_B1.py <ruta_al_xlsx>")
        sys.exit(1)

    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    ws_tramos = wb["Tramos"]
    ws_dif = wb["Diferencias"]

    tramos = []
    for r in range(2, ws_tramos.max_row + 1):
        entr = ws_tramos.cell(r, 1).value
        if not entr:
            continue
        desde = ws_tramos.cell(r, 2).value
        hasta = ws_tramos.cell(r, 3).value
        palabras = ws_tramos.cell(r, 4).value
        tramos.append({
            "entr": entr, "desde": desde, "hasta": hasta,
            "palabras": palabras, "clave": f"{entr} ({desde}-{hasta})",
            "diferencias": 0, "no_verificadas": [],
        })

    cache_transcripciones = {}

    total_filas = 0
    total_no_verificadas = 0
    for r in range(2, ws_dif.max_row + 1):
        entr = ws_dif.cell(r, 1).value
        minuto = ws_dif.cell(r, 2).value
        escrito = ws_dif.cell(r, 3).value
        oido = ws_dif.cell(r, 4).value
        if not entr or not minuto or not escrito:
            continue
        total_filas += 1

        if entr not in cache_transcripciones:
            cache_transcripciones[entr] = cargar_transcripcion(entr)
        bloques = cache_transcripciones[entr]

        seg = minuto_a_segundos(minuto)
        # buscar el bloque cuyo timestamp esta MAS CERCA del minuto indicado
        candidato = None
        mejor_diff = None
        for ts, contenido in bloques:
            diff = abs(ts - seg)
            if diff <= 20 and (mejor_diff is None or diff < mejor_diff):
                mejor_diff = diff
                candidato = contenido

        ok = False
        if candidato and normalizar(escrito) in normalizar(candidato):
            ok = True

        # encontrar a que tramo pertenece esta fila
        tramo_match = None
        for t in tramos:
            if t["entr"] != entr:
                continue
            d, h = minuto_a_segundos(t["desde"]), minuto_a_segundos(t["hasta"])
            if d - 5 <= seg <= h + 5:
                tramo_match = t
                break

        if ok:
            if tramo_match:
                tramo_match["diferencias"] += 1
        else:
            total_no_verificadas += 1
            msg = f"{entr} {minuto}: NO SE ENCONTRO el texto 'escrito' en la transcripcion real"
            if tramo_match:
                tramo_match["no_verificadas"].append(msg)
            print("[FALLA]", msg)
            print("   escrito declarado:", repr(escrito))
            if candidato:
                print("   contenido real mas cercano:", repr(candidato[:200]))

    print()
    print("=" * 70)
    print(f"Filas totales en 'Diferencias': {total_filas}")
    print(f"Filas que NO se pudieron verificar contra el archivo real: {total_no_verificadas}")
    print("=" * 70)
    print()
    print(f"{'Tramo':40s} {'palabras':>9s} {'dif.reales':>11s} {'max 5%':>8s} {'%':>7s} {'cumple':>7s}")
    for t in tramos:
        max5 = round(t["palabras"] * 0.05)
        pct = 100 * t["diferencias"] / t["palabras"] if t["palabras"] else 0
        cumple = "SI" if t["diferencias"] <= max5 else "NO"
        print(f"{t['clave']:40s} {t['palabras']:9d} {t['diferencias']:11d} {max5:8d} {pct:6.2f}% {cumple:>7s}")
        if t["no_verificadas"]:
            for m in t["no_verificadas"]:
                print("    ! ", m)


if __name__ == "__main__":
    main()
