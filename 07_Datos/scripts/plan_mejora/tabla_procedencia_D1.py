#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
D1 - Tabla de procedencia de TODOS los requisitos (RF, RNF, RF-IA y RNF-IA).

Deriva del propio repositorio, sin inventar nada:
  - el catalogo completo de requisitos y el origen que cada uno declara;
  - la fecha y el tipo de cada evidencia, tomados del apendice del ERS;
  - el commit en que cada identificador aparece por primera vez (git log -S);
  - una CLASIFICACION PROPUESTA con el motivo que la sustenta.

La clasificacion es una PROPUESTA derivada de reglas explicitas. La columna
CLASIFICACION_VERIFICADA la rellena una persona. Nada se marca como
verificado por el script.

Uso:
    python3 tabla_procedencia_D1.py [--fusionar <csv_o_xlsx_ya_revisado>]

Con --fusionar se conservan las columnas CLASIFICACION_VERIFICADA y
VERIFICADO_POR ya rellenadas por una persona, emparejadas por id_requisito.

---------------------------------------------------------------------------
Historial de correcciones de este script (tarea D1 del Plan de mejora)
---------------------------------------------------------------------------
v2 (20/09/2026) - correcciones sobre la primera version:

  [1] Extraccion del origen de los RF.  La version anterior usaba
      \\RF\\{(RF-\\d+)\\}\\{([^}]*)\\}\\s*\\{(.*?)\\}\\s*\\{(.*?)\\}
      El cierre no goloso `(.*?)\\}` paraba en la primera llave de cierre,
      que pertenece al `\\id{...}` anidado dentro del argumento. Resultado:
      el origen salia cortado en las 42 filas de RF
      (p. ej. RF-01 -> "Administrador general . \\id{EV-01").
      Se sustituye por un lector de argumentos con balance de llaves.
      Se elimina ademas el truncado [:300] al escribir el CSV.

  [2] Deteccion de requisitos derivados.  La version anterior usaba
      (?:RF|RNF|RD|L)-(?:IA-)?\\d+
      La alternativa "L" suelta y la ausencia de limite por la izquierda
      hacian que "RL-04" (requisito legal, Art. 12 LOPDP) casara como
      "L-04", y RF-40, RF-41 y RF-42 quedaran clasificados "derivado"
      cuando su origen es normativo. Se anade limite por la izquierda y se
      separan los identificadores legales (RL-XX, L-XX) de los de
      requisito, tratandolos como senal normativa.

  [3] Catalogo incompleto.  La version anterior recogia 42 RF + 19 RNF +
      12 RNF-IA = 73 requisitos, pero el ERS declara "dieciocho requisitos
      de inteligencia artificial": faltaban los 6 RF-IA de la seccion 9
      (lineas 85-87 y 149-151 de seccion9_ia.tex). El criterio de D1 pide
      el 100 % de los requisitos de IA. El catalogo pasa a 79.

  [4] Origen de los requisitos de IA.  Su tabla no tiene columna de origen
      (RF-IA: 2 columnas; RNF-IA: 3 columnas), de modo que la version
      anterior tomaba como "origen" el texto completo del requisito. Eso
      convertia cualquier referencia incidental en una derivacion. Ahora el
      origen de los requisitos de IA se toma de la tabla de trazabilidad
      del propio ERS (seccion9_ia.tex, tab:traza-ia), que declara de que
      requisito preexistente deriva cada uno, con expansion de los rangos
      del tipo "RNF-IA-01 a RNF-IA-03".

  [5] Commit de alta.  La version anterior buscaba el identificador con
      `git log -S <id> -- 01_ERS/ERS_SRS_2B_v2.0.tex` (solo la ruta actual)
      y devolvia el commit 2624c0f (03/09, migracion del repositorio) para los
      79 requisitos, con lo que ninguna alerta de anterioridad podia
      dispararse. El ERS existia antes como AHMRV/01_ERS/ERS_SRS_2A_v1.0.tex.
      Ahora se recorre en una sola pasada el historial de todas las versiones
      .tex (`-- '*01_ERS/*.tex'`) y se toma el primer commit que anade cada
      identificador: 53 requisitos el 02/08 (35fff3c), RF-36 a RF-39 el 02/08
      (54bd614) y los 22 restantes el 31/08 (14b9e6b).
      Ademas, para las evidencias con rango de fechas ("03--07/08/2026") la
      fecha se toma por el INICIO del rango (antes tomaba el final).
---------------------------------------------------------------------------
"""
import csv, re, subprocess, os, sys, collections

ERS       = "01_ERS/ERS_SRS_2B_v2.0.tex"
IA        = "01_ERS/seccion9_ia.tex"
APEND     = "01_ERS/apendices.tex"
PROPUESTA = "01_ERS/antecedentes/2026-05-05_Propuesta_Inicial_1A.pdf"
SALIDA    = "07_Datos/datos_procesados/tabla_procedencia_requisitos.csv"

FECHA_PROPUESTA = "2026-05-05"

# Requisitos que el Plan de mejora de datos (19/09/2026) senala expresamente
# como ya presentes en la propuesta del 05/05, antes de la primera entrevista.
SENALADOS_DOCENTE = {"RF-03", "RF-07", "RF-08", "RF-18"}


def leer(p):
    return open(p, encoding="utf-8").read()


# =========================================================================
#  Utilidades de LaTeX
# =========================================================================

def arg_balanceado(texto, i):
    """Lee un argumento {...} que empieza en texto[i] == '{'.

    Cuenta llaves de apertura y cierre para no pararse en las llaves de un
    \\id{...} anidado. Devuelve (contenido, indice_tras_la_llave_de_cierre).
    """
    if i >= len(texto) or texto[i] != "{":
        return None, i
    prof, j = 0, i
    while j < len(texto):
        c = texto[j]
        if c == "\\":                 # \{ y \} son literales, no delimitan
            j += 2
            continue
        if c == "{":
            prof += 1
        elif c == "}":
            prof -= 1
            if prof == 0:
                return texto[i + 1:j], j + 1
        j += 1
    return None, i


def args_macro(texto, pos, n):
    """Lee n argumentos consecutivos {..}{..}... a partir de pos."""
    out = []
    i = pos
    for _ in range(n):
        while i < len(texto) and texto[i] in " \t\r\n%":
            if texto[i] == "%":       # comentario de LaTeX
                while i < len(texto) and texto[i] != "\n":
                    i += 1
            i += 1
        a, i = arg_balanceado(texto, i)
        if a is None:
            return None, pos
        out.append(a)
    return out, i


def celdas(fila):
    """Parte una fila de tabla por & sin romper en \\& ni dentro de {...}."""
    out, buf, prof, i = [], [], 0, 0
    while i < len(fila):
        c = fila[i]
        if c == "\\" and i + 1 < len(fila):
            buf.append(fila[i:i + 2]); i += 2; continue
        if c == "{":
            prof += 1
        elif c == "}":
            prof -= 1
        if c == "&" and prof == 0:
            out.append("".join(buf)); buf = []
        else:
            buf.append(c)
        i += 1
    out.append("".join(buf))
    return out


def limpiar(t):
    return re.sub(r"\s+", " ", (t or "")).strip()


# =========================================================================
#  Identificadores
#  RF-XX / RNF-XX / RD-XX / RF-IA-XX / RNF-IA-XX -> requisitos
#  RL-XX / L-XX                                  -> legal y limitaciones
#  El limite por la izquierda impide que "RL-04" case como "L-04".
# =========================================================================
RE_REQ   = re.compile(r"(?<![A-Za-z])(?:RF-IA|RNF-IA|RF|RNF|RD)-\d+(?![\w-])")
RE_LEGAL = re.compile(r"(?<![A-Za-z])(?:RL|L)-\d+(?![\w-])")
RE_EV    = re.compile(r"(?<![A-Za-z])EV-\d+(?![\w-])")

# Con limites de palabra: sin ellos, "disociada" contenia "iso" y clasificaba
# RF-IA-03 como normativo por una coincidencia dentro de una palabra.
NORMA = re.compile(r"\b(?:LOPDP|ISO|IEC|25010|29148)\b|\blegal|\bnormativ|\bArt\.\s*\d+", re.I)

# El ERS usa la palabra "Derivado" como valor literal de la columna Origen de
# algunos RNF, sin nombrar el requisito del que derivan.
DERIV_LIT = re.compile(r"\bderivad[oa]s?\b", re.I)


def orden_id(rid):
    m = re.search(r"(\d+)$", rid)
    return int(m.group(1)) if m else 0


# =========================================================================
#  Fechas y tipo de cada evidencia (apendice del ERS)
# =========================================================================
FECHA_EV = {}
for m in re.finditer(r"\\id\{(EV-\d+)\}\s*&\s*([^&]+?)\s*&\s*([^&]+?)\s*&", leer(APEND)):
    ev, tipo, fecha = m.group(1), m.group(2).strip(), m.group(3).strip()
    if ev not in FECHA_EV:
        FECHA_EV[ev] = (tipo, fecha)


def fecha_iso(txt):
    """'23/05/2026' -> 2026-05-23. Para rangos '03--07/08/2026' devuelve el INICIO (2026-08-03)."""
    m = re.search(r"(\d{1,2})(?:\s*--\s*\d{1,2})?/(\d{2})/(\d{4})", txt or "")
    return f"{m.group(3)}-{m.group(2)}-{int(m.group(1)):02d}" if m else ""


# =========================================================================
#  Modulos de la propuesta inicial (05/05/2026)
#  Correspondencia PROPUESTA por coincidencia de terminos, no un juicio.
# =========================================================================
MODULOS = []
try:
    _t = subprocess.run(["pdftotext", "-layout", PROPUESTA, "-"],
                        capture_output=True, text=True, timeout=60).stdout
    for _s in re.split(r"\n(?=\d+\.\d+\s)", _t):
        _m = re.match(r"(\d+\.\d+)\s+(.+)", _s)
        if _m:
            MODULOS.append((_m.group(1), _m.group(2).strip(),
                            re.sub(r"\s+", " ", _s).lower()))
except Exception:
    pass

_STOP = set("de la el en y a los las del que se un una por con para es su al lo como mas o "
            "sistema debe permitir gestion registro control".split())


def en_propuesta(nombre):
    """Correspondencia PROPUESTA entre el nombre del requisito y un modulo de
    la propuesta del 05/05. Los requisitos de IA no tienen nombre corto sino
    un enunciado completo, y sobre un texto largo casi cualquier modulo
    acumula coincidencias; por eso el umbral sube con la longitud."""
    texto = nombre[:140]
    pal = {w for w in re.sub(r"[^a-zA-ZáéíóúñÁÉÍÓÚÑ ]", " ", texto.lower()).split()
           if len(w) > 4 and w not in _STOP}
    if not pal:
        return ""
    umbral = 2 if len(texto) <= 80 else 4
    mejor = (0, "")
    for num, tit, cuerpo in MODULOS:
        n = sum(1 for w in pal if w[:6] in cuerpo)
        if n > mejor[0]:
            mejor = (n, f"{num} {tit}")
    return mejor[1] if mejor[0] >= umbral else ""


# =========================================================================
#  Commit de alta de cada identificador
# =========================================================================
CACHE = {}
_ALTA_LEIDA = False


def _leer_historial_alta():
    """
    Una sola pasada por el historial de TODAS las versiones .tex del ERS
    (pathspec *01_ERS/*.tex: incluye AHMRV/01_ERS/ERS_SRS_2A_v1.0.tex, anterior
    a la migracion del 03/09) y guarda, para cada identificador, el primer
    commit cuya diferencia lo AÑADE.
    """
    global _ALTA_LEIDA
    _ALTA_LEIDA = True
    try:
        out = subprocess.run(
            ["git", "log", "--reverse", "--format=COMMIT %h %ad", "--date=short",
             "-p", "--unified=0", "--", "*01_ERS/*.tex"],
            capture_output=True, text=True, errors="ignore", timeout=180).stdout
    except Exception:
        return
    actual = ("", "")
    for linea in out.splitlines():
        if linea.startswith("COMMIT "):
            _, h, d = linea.split()[:3]
            actual = (h, d)
        elif linea.startswith("+") and not linea.startswith("+++"):
            for rid in RE_REQ.findall(linea):
                CACHE.setdefault(rid, "%s|%s" % actual)


def commit_alta(rid, ficheros):
    """Primer commit del historial completo que anade el identificador ('sha|fecha')."""
    if not _ALTA_LEIDA:
        _leer_historial_alta()
    return CACHE.get(rid, "|")


# =========================================================================
#  Tabla de trazabilidad de los requisitos de IA (seccion9_ia.tex)
#  Su tabla de requisitos no tiene columna de origen; esta si lo declara.
# =========================================================================
t_ia = leer(IA)
TRAZA_IA = {}

_bloque = re.search(r"Requisito IA.*?\\end\{longtable\}", t_ia, re.S)
if _bloque:
    for _linea in _bloque.group(0).splitlines():
        if "\\id{" not in _linea or "&" not in _linea:
            continue
        cs = celdas(_linea)
        if len(cs) < 5:
            continue
        col_ids, col_orig, col_rel = cs[0], cs[1], cs[4]
        ids = RE_REQ.findall(col_ids)
        # "RNF-IA-01 a RNF-IA-03" -> expandir el rango
        if len(ids) == 2 and re.search(r"\}\s*a\s*\\id", col_ids):
            pre = ids[0].rsplit("-", 1)[0]
            if ids[1].rsplit("-", 1)[0] == pre:
                ids = [f"{pre}-{n:02d}" for n in range(orden_id(ids[0]), orden_id(ids[1]) + 1)]
        rel = limpiar(re.sub(r"\\\\\s*\\hline\s*$", "", col_rel))
        for rid in ids:
            TRAZA_IA[rid] = (limpiar(col_orig), rel)


# =========================================================================
#  Catalogo de requisitos
# =========================================================================
reqs = []
t_ers = leer(ERS)

# --- RF: macro \RF{ID}{Nombre}{Descripcion}{Actor/Origen}{...}{...}{...}{...}
#     Lectura con balance de llaves: el 4.o argumento contiene \id{EV-XX}.
for m in re.finditer(r"\\RF\s*(?=\{)", t_ers):
    a, _ = args_macro(t_ers, m.end(), 8)
    if not a:
        continue
    rid = limpiar(a[0])
    if not re.fullmatch(r"RF-\d+", rid):
        continue
    reqs.append({"id": rid, "tipo": "RF", "nombre": limpiar(a[1]),
                 "origen_ers": limpiar(a[3]),
                 "origen_fuente": "campo Actor / Origen (evidencia) del ERS",
                 "fuente": ERS})

# --- RNF: fila de 5 columnas, la ultima es Origen
for m in re.finditer(r"\\id\{(RNF-\d+)\}\s*&(.*?)\\\\ *\\hline", t_ers, re.S):
    cs = celdas(m.group(2))
    if len(cs) < 4:
        continue
    reqs.append({"id": m.group(1), "tipo": "RNF", "nombre": limpiar(cs[0]),
                 "origen_ers": limpiar(cs[-1]),
                 "origen_fuente": "columna Origen de la tabla de RNF del ERS",
                 "fuente": ERS})

# --- RF-IA: fila de 2 columnas (ID | requisito). Sin columna de origen.
for m in re.finditer(r"\\id\{(RF-IA-\d+)\}\s*&(.*?)\\\\ *\\hline", t_ia, re.S):
    cs = celdas(m.group(2))
    if len(cs) != 1:            # las filas de la tabla de trazabilidad tienen 5
        continue
    texto = limpiar(cs[0])
    nombre = limpiar(re.split(r"\\newline|\\emph\{Criterio", texto)[0])[:200]
    orig, rel = TRAZA_IA.get(m.group(1), ("", ""))
    reqs.append({"id": m.group(1), "tipo": "RF-IA", "nombre": nombre,
                 "origen_ers": (f"{orig} - {rel}" if orig else ""),
                 "origen_fuente": "tabla de trazabilidad de los requisitos de IA (tab:traza-ia)",
                 "fuente": IA, "texto_requisito": texto})

# --- RNF-IA: fila de 3 columnas (ID | Tipo | requisito). Sin columna de origen.
for m in re.finditer(r"\\id\{(RNF-IA-\d+)\}\s*&(.*?)\\\\ *\\hline", t_ia, re.S):
    cs = celdas(m.group(2))
    if len(cs) != 2:
        continue
    texto = limpiar(cs[1])
    orig, rel = TRAZA_IA.get(m.group(1), ("", ""))
    reqs.append({"id": m.group(1), "tipo": "RNF-IA", "nombre": limpiar(cs[0]),
                 "origen_ers": (f"{orig} - {rel}" if orig else ""),
                 "origen_fuente": "tabla de trazabilidad de los requisitos de IA (tab:traza-ia)",
                 "fuente": IA, "texto_requisito": texto})

vistos, uniq = set(), []
for r in reqs:
    if r["id"] not in vistos:
        vistos.add(r["id"])
        uniq.append(r)
reqs = uniq


# =========================================================================
#  Clasificacion propuesta
# =========================================================================
ORDEN_TIPO = {"RF": 0, "RF-IA": 1, "RNF": 2, "RNF-IA": 3}
filas = []

for r in reqs:
    origen = r["origen_ers"]

    evs    = sorted(set(RE_EV.findall(origen)), key=orden_id)
    refs   = sorted(set(RE_REQ.findall(origen)) - {r["id"]},
                    key=lambda x: (x.rsplit("-", 1)[0], orden_id(x)))
    leg    = sorted(set(RE_LEGAL.findall(origen)), key=orden_id)

    sha, falta = (commit_alta(r["id"], [r["fuente"]]).split("|") + [""])[:2]

    fechas = [fecha_iso(FECHA_EV.get(e, ("", ""))[1]) for e in evs]
    fechas = [f for f in fechas if f]
    f_ev = min(fechas) if fechas else ""

    origen_sin_macro = re.sub(r"\\[a-zA-Z]+", " ", origen)   # fuera \footnotesize, \id, \cite

    if evs:
        clas, motivo = "elicitado", f"cita evidencia: {', '.join(evs)}"
    elif any(x.startswith("RL-") for x in leg) or NORMA.search(origen + " " + r["nombre"]):
        ref_leg = ", ".join(x for x in leg if x.startswith("RL-"))
        motivo = ("el origen declarado invoca requisito legal: " + ref_leg) if ref_leg \
                 else "el origen declarado invoca norma o marco legal"
        clas = "normativo"
    elif refs:
        clas, motivo = "derivado", f"se apoya en requisito(s) previo(s): {', '.join(refs)}"
    elif DERIV_LIT.search(origen_sin_macro):
        clas = "derivado"
        motivo = ("el ERS declara el origen literalmente como «Derivado», "
                  "sin nombrar el requisito de procedencia")
    elif origen_sin_macro.strip():
        clas, motivo = "propuesta_equipo", "declara origen, pero sin evidencia, norma ni requisito previo"
    else:
        clas, motivo = "propuesta_equipo", "no declara origen alguno en el ERS"

    modulo = en_propuesta(r["nombre"])
    if r["id"] in SENALADOS_DOCENTE and not modulo:
        modulo = "senalado en el plan de mejora (correspondencia a verificar)"

    alerta = []
    if modulo and f_ev and FECHA_PROPUESTA < f_ev:
        alerta.append(f"FUNCIONALIDAD YA EN LA PROPUESTA DEL {FECHA_PROPUESTA} "
                      f"({modulo}), anterior a su evidencia {f_ev}")
    if f_ev and falta and falta < f_ev:
        alerta.append(f"REQUISITO ANTERIOR A SU EVIDENCIA (alta {falta} < evidencia {f_ev})")
    if not evs and r["tipo"] == "RF":
        alerta.append("RF sin evidencia citada")
    if not evs and r["tipo"] in ("RF-IA", "RNF-IA"):
        alerta.append("requisito de IA sin evidencia de campo; su origen es la tabla de "
                      "trazabilidad, no una entrevista")
    if clas == "derivado" and not refs:
        alerta.append("ORIGEN «Derivado» SIN REQUISITO DE PROCEDENCIA IDENTIFICADO")
    if not origen:
        alerta.append("SIN ORIGEN DECLARADO EN EL ERS")

    filas.append({
        "id_requisito": r["id"],
        "tipo": r["tipo"],
        "nombre": r["nombre"][:120],
        "evidencias_citadas": "; ".join(evs),
        "fecha_evidencia_mas_antigua": f_ev,
        "tipo_evidencia": "; ".join(FECHA_EV.get(e, ("", ""))[0] for e in evs),
        "requisitos_referidos": "; ".join(refs),
        "referencias_legales": "; ".join(leg),
        "commit_alta": sha,
        "fecha_commit_alta": falta,
        "CLASIFICACION_PROPUESTA": clas,
        "MOTIVO_PROPUESTA": motivo,
        "en_propuesta_inicial": modulo,
        "CLASIFICACION_VERIFICADA": "",
        "VERIFICADO_POR": "",
        "ALERTA": " | ".join(alerta),
        "origen_completo_ers": origen,
        "de_donde_sale_el_origen": r["origen_fuente"],
    })

filas.sort(key=lambda x: (ORDEN_TIPO.get(x["tipo"], 9), orden_id(x["id_requisito"])))


# =========================================================================
#  Fusion opcional con una hoja ya revisada por una persona
# =========================================================================
if "--fusionar" in sys.argv:
    prev = sys.argv[sys.argv.index("--fusionar") + 1]
    hechas = {}
    def _norm(s):
        return re.sub(r"[^a-z]", "", str(s or "").lower())

    if prev.lower().endswith(".xlsx"):
        import openpyxl
        ws = openpyxl.load_workbook(prev, data_only=True).active
        # La hoja usa titulos legibles ("CLASIFICACION VERIFICADA"); el CSV usa
        # claves con guion bajo. Se emparejan normalizando el titulo.
        ic = {_norm(c.value): i for i, c in enumerate(ws[1]) if c.value}

        def col(*alias):
            for a in alias:
                if a in ic:
                    return ic[a]
            sys.exit(f"La hoja {prev} no tiene ninguna columna {alias}")

        i_id  = col("idrequisito", "id")
        i_cla = col("clasificacionverificada")
        i_ver = col("verificadopor")
        for fila in ws.iter_rows(min_row=2, values_only=True):
            rid = fila[i_id]
            if rid:
                hechas[str(rid).strip()] = (fila[i_cla] or "", fila[i_ver] or "")
    else:
        with open(prev, encoding="utf-8-sig", newline="") as f:
            for d in csv.DictReader(f, delimiter=";"):
                hechas[d["id_requisito"]] = (d.get("CLASIFICACION_VERIFICADA", ""),
                                             d.get("VERIFICADO_POR", ""))
    n = 0
    for x in filas:
        cv, vp = hechas.get(x["id_requisito"], ("", ""))
        # Solo cuenta como verificada una fila con clasificacion Y firma de una
        # persona: una clasificacion prellenada sin firma NO es una verificacion.
        if str(cv).strip() and str(vp).strip():
            x["CLASIFICACION_VERIFICADA"], x["VERIFICADO_POR"] = str(cv).strip(), str(vp).strip()
            n += 1
    print(f"Fusionadas {n} filas ya verificadas desde {prev}")


# =========================================================================
#  Salida
# =========================================================================
os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
with open(SALIDA, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(filas[0].keys()), delimiter=";")
    w.writeheader()
    w.writerows(filas)

print(f"Requisitos: {len(filas)}  ->  {dict(collections.Counter(x['tipo'] for x in filas))}")
print(f"Clasificacion propuesta: {dict(collections.Counter(x['CLASIFICACION_PROPUESTA'] for x in filas))}")
print(f"En la propuesta del 05/05: {sum(1 for x in filas if x['en_propuesta_inicial'])}")
print(f"Verificadas por persona: {sum(1 for x in filas if x['CLASIFICACION_VERIFICADA'])} de {len(filas)}")
al = [x for x in filas if x["ALERTA"]]
print(f"Con alerta: {len(al)}")
for x in al:
    print(f"   {x['id_requisito']:11s} {x['ALERTA'][:96]}")
print(f"\nEscrito: {SALIDA}")
