#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
preparar_muestra_C3.py  ·  C3 · SIMPA_ISR401

Congela la muestra de la doble codificación REAL (tarea C3 del plan de mejora).

Criterio de C3: "al menos el 20 % de los fragmentos elegidos al azar con semilla
declarada; cada codificador sube su archivo desde su propia cuenta sin ver el del otro;
libro de códigos versionado antes de los dos archivos; kappa por script y tabla de
desacuerdos".

QUÉ HACE
    · Toma los fragmentos codificados de los dos CSV (dominio y contraste) y elige
      TAMANO al azar, sin reemplazo, con la semilla SEMILLA. Es determinista: quien ejecute
      este script obtiene exactamente la misma muestra (así se puede auditar).
    · Escribe muestra_C3_fragmentos.csv y la hoja para codificar hoja_codificacion_C3.xlsx
      (lista desplegable con los 95 códigos del libro de códigos v1.0 y sus definiciones).
    · NO incluye el código original en ninguna hoja de los codificadores.

USO (desde la raíz del repositorio)
    python3 10_Autoria/doble_codificacion_C3/preparar_muestra_C3.py
No volver a ejecutar después de que alguien empiece a codificar (cambiaría los archivos).
"""
import csv
import random
import re
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[2]
SALIDA = Path(__file__).resolve().parent
FUENTES = [RAIZ / "07_Datos" / "datos_crudos" / "codificacion.csv",
           RAIZ / "07_Datos" / "datos_procesados" / "codificacion_tercera_ronda.csv"]
LIBRO = RAIZ / "07_Datos" / "libro_codigos.md"
SEMILLA = 20260921
TAMANO = 45          # 45 de 213 fragmentos = 21,1 % (criterio: al menos 20 %)


def cargar_fragmentos():
    filas = []
    for ruta in FUENTES:
        with open(ruta, newline="", encoding="utf-8-sig") as f:
            for r in csv.DictReader(f, delimiter=";"):
                r["_origen"] = ruta.name
                filas.append(r)
    return filas


def cargar_libro():
    codigos = []
    for linea in LIBRO.read_text(encoding="utf-8").splitlines():
        m = re.match(r"\|\s*(\d+)\s*\|\s*([A-ZÑ0-9_]+)\s*\|\s*(.+?)\s*\|\s*(.+?)\s*\|\s*$", linea)
        if m:
            codigos.append((int(m.group(1)), m.group(2), m.group(3), m.group(4)))
    return codigos


def muestra():
    filas = cargar_fragmentos()
    elegidos = random.Random(SEMILLA).sample(range(len(filas)), TAMANO)
    return [(f"C3-{i:03d}", filas[k]) for i, k in enumerate(elegidos, start=1)]


def muestra_congelada():
    """Lee muestra_C3_fragmentos.csv (la muestra vigente) y devuelve [(ID, fila original con su código)]."""
    orig = {(r["_origen"], r["linea_csv"]): r for r in cargar_fragmentos()}
    with open(SALIDA / "muestra_C3_fragmentos.csv", newline="", encoding="utf-8-sig") as f:
        return [(m["ID_fragmento"], orig[(m["archivo_origen"], m["linea_csv"])]) for m in csv.DictReader(f, delimiter=";")]


def main():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    sel = muestra()
    libro = cargar_libro()
    with open(SALIDA / "muestra_C3_fragmentos.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["ID_fragmento", "Fragmento", "Cita_literal", "Entrevista", "archivo_origen", "linea_csv"])
        for id_, r in sel:
            w.writerow([id_, r["Fragmento_parafraseado"], r["CITA_LITERAL"], re.search(r"ENTR-\d+", r["transcripcion"]).group(0), r["_origen"], r["linea_csv"]])

    fuente = "Arial"
    amarillo = PatternFill("solid", start_color="FFF2CC")
    azul = PatternFill("solid", start_color="1F3864")
    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"
    texto = [
        ("Doble codificación C3: instrucciones", True),
        ("1. Codifica SOLO con esta hoja. No abras codificacion.csv ni el archivo de otra persona.", False),
        ("2. En la hoja 'Codificar', para cada fragmento elige en la columna amarilla CODIGO uno de los 95 códigos de la lista.", False),
        ("3. Si dudas, consulta la hoja 'Libro de codigos' (definición y criterio de cada código). Escoge UN solo código por fragmento.", False),
        ("4. Escribe tu nombre completo en la columna CODIFICADOR (en todas las filas) y, si quieres, una nota en OBSERVACION.", False),
        ("5. Guarda el archivo como  codificacion_C3_TUAPELLIDO.xlsx  y súbelo TÚ MISMO desde TU cuenta de GitHub a 10_Autoria/doble_codificacion_C3/.", False),
        ("6. No modifiques los IDs, los fragmentos ni las hojas. No edites el archivo de otra persona.", False),
        ("Muestra: 45 fragmentos de 213 (21,1 %), elegidos al azar con la semilla 20260921 (ver preparar_muestra_C3.py).", False),
    ]
    for t, negrita in texto:
        ins.append([t])
        ins.cell(row=ins.max_row, column=1).font = Font(name=fuente, size=10, bold=negrita)
    ins.column_dimensions["A"].width = 135

    ws = wb.create_sheet("Codificar")
    ws.append(["ID_fragmento", "Fragmento", "Cita_literal", "Entrevista", "CODIGO", "CODIFICADOR", "OBSERVACION"])
    for id_, r in sel:
        ws.append([id_, r["Fragmento_parafraseado"], r["CITA_LITERAL"],
                   re.search(r"ENTR-\d+", r["transcripcion"]).group(0), None, None, None])
    for c, ancho in enumerate([12, 60, 70, 11, 30, 22, 36], start=1):
        ws.column_dimensions[get_column_letter(c)].width = ancho
        h = ws.cell(row=1, column=c)
        h.font = Font(name=fuente, bold=True, color="FFFFFF", size=10)
        h.fill = azul
        h.alignment = Alignment(wrap_text=True, vertical="center")
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.font = Font(name=fuente, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.column >= 5:
                c.fill = amarillo
    dv = DataValidation(type="list", formula1=f"='Libro de codigos'!$B$2:$B${len(libro) + 1}", allow_blank=True)
    dv.error, dv.errorTitle, dv.showErrorMessage = "Elige un código de la lista.", "Código no válido", True
    ws.add_data_validation(dv)
    dv.add(f"E2:E{len(sel) + 1}")
    ws.freeze_panes = "C2"

    wl = wb.create_sheet("Libro de codigos")
    wl.append(["#", "Codigo", "Definicion", "Criterio de aplicacion"])
    for fila in libro:
        wl.append(list(fila))
    for c, ancho in enumerate([5, 30, 70, 70], start=1):
        wl.column_dimensions[get_column_letter(c)].width = ancho
        wl.cell(row=1, column=c).font = Font(name=fuente, bold=True)
    for fila in wl.iter_rows(min_row=2):
        for c in fila:
            c.font = Font(name=fuente, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(SALIDA / "hoja_codificacion_C3.xlsx")
    print(f"Muestra congelada: {len(sel)} de {len(cargar_fragmentos())} fragmentos, semilla {SEMILLA}; libro con {len(libro)} códigos.")


if __name__ == "__main__":
    main()
