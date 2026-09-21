#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
reanclar_citas_C1.py  ·  C1 (y de paso C2) · SIMPA_ISR401

PROBLEMA
    Tras la retranscripción literal de B1 (ENTR-01 a 08), las citas de
    codificacion.csv dejaron de ser literales y sus números de línea ya no
    corresponden. Además el analizador antiguo no reconocía los nuevos formatos
    de hablante (ver citas_lib.py). Quedan ~153 fragmentos sin cita válida.

QUÉ HACE ESTE SCRIPT
    1) proponer  -> genera una hoja de revisión (.xlsx) SOLO con los fragmentos
                    cuya cita hoy no es literal o no existe. Para cada uno ofrece
                    3 candidatas (subcadenas EXACTAS del entrevistado) y avisa si
                    otra entrevista respalda mejor el fragmento (insumo de C2).
    2) aplicar   -> lee la hoja ya revisada por personas y actualiza los CSV de
                    codificación SOLO en las filas con decisión, aceptando únicamente
                    citas que se comprueben literales. Las filas sin decisión y las
                    ya verificadas NO se tocan.

    Ninguna decisión la toma el script: una persona confirma que la cita
    RESPALDA el código (el verificador solo comprueba que el texto exista).

    OJO: no usar aplicar_revision_C1.py con una hoja parcial; ese script deja
    en PENDIENTE_REVISION todas las filas que no aparecen en la hoja.

USO (desde la raíz del repositorio)
    python3 07_Datos/scripts/plan_mejora/reanclar_citas_C1.py proponer
    python3 07_Datos/scripts/plan_mejora/reanclar_citas_C1.py aplicar HOJA.xlsx --simular
    python3 07_Datos/scripts/plan_mejora/reanclar_citas_C1.py aplicar HOJA.xlsx
    python3 07_Datos/scripts/plan_mejora/verificar_citas_C1.py      # control final
"""
import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import citas_lib as L  # noqa: E402

RAIZ = L.RAIZ
ARCHIVOS_CSV = {
    "codificacion.csv": RAIZ / "07_Datos" / "datos_crudos" / "codificacion.csv",
    "codificacion_tercera_ronda.csv": RAIZ / "07_Datos" / "datos_procesados" / "codificacion_tercera_ronda.csv",
}
HOJA_SALIDA = RAIZ / "07_Datos" / "resultados" / "revision_citas_C1_reanclaje.xlsx"
UMBRAL_OTRA, MARGEN_OTRA = 0.40, 0.10   # alerta "mejor respaldo en otra entrevista"
DECISIONES = ["CAND1", "CAND2", "CAND3", "LITERAL", "NO_LOCALIZADA"]


# ------------------------------------------------------------------ CSV
def leer_csv(ruta):
    b = ruta.read_bytes()
    bom, crlf = b.startswith(b"\xef\xbb\xbf"), b"\r\n" in b[:3000]
    with open(ruta, newline="", encoding="utf-8-sig") as f:
        lector = csv.DictReader(f, delimiter=";")
        filas = list(lector)
        return filas, list(lector.fieldnames), bom, crlf


def escribir_csv(ruta, filas, campos, bom, crlf):
    with open(ruta, "w", newline="", encoding="utf-8-sig" if bom else "utf-8") as f:
        w = csv.DictWriter(f, fieldnames=campos, delimiter=";",
                           lineterminator="\r\n" if crlf else "\n")
        w.writeheader()
        w.writerows(filas)


def cita_ok(fila, turnos):
    n = fila.get("LINEA_TRANSCRIPCION", "").strip()
    if fila.get("ESTADO", "").strip() != "VER" or not n.isdigit():
        return False
    return any(nn == int(n) and fila["CITA_LITERAL"] in l[o:] for nn, l, o in turnos[fila["transcripcion"]])


def cargar_turnos():
    return {f.name: L.turnos_de_archivo(f)
            for f in sorted(L.TRANSCRIPCIONES.glob("2026-*_Transcripcion.md"))}


def corto(nombre):
    m = re.search(r"ENTR-\d+", nombre)
    return m.group(0) if m else nombre


# -------------------------------------------------------------- proponer
def proponer():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    turnos = cargar_turnos()
    registros = []
    for nombre, ruta in ARCHIVOS_CSV.items():
        filas, *_ = leer_csv(ruta)
        for fila in filas:
            if cita_ok(fila, turnos):
                continue
            propia = L.candidatas(turnos[fila["transcripcion"]], fila["Fragmento_parafraseado"],
                                  fila["CITA_LITERAL"] if fila["ESTADO"].strip() == "VER" else "", k=3)
            mejor = (0.0, None, None, None)
            for otra, tr in turnos.items():
                if otra == fila["transcripcion"]:
                    continue
                c = L.candidatas(tr, fila["Fragmento_parafraseado"], "", k=1)
                if c and c[0][0] > mejor[0]:
                    mejor = (c[0][0], otra, c[0][1], c[0][2])
            p1 = propia[0][0] if propia else 0.0
            alerta = ""
            if mejor[0] >= UMBRAL_OTRA and mejor[0] >= p1 + MARGEN_OTRA:
                alerta = (f"{corto(mejor[1])} línea {mejor[2]} respalda mejor ({mejor[0]:.2f} vs {p1:.2f}): "
                          f"«{mejor[3][:160]}»")
            registros.append((nombre, fila, propia, alerta))

    wb = Workbook()
    ws = wb.active
    ws.title = "Revision"
    cab = ["N", "Archivo_CSV", "Fila_CSV", "ID_evidencia", "Transcripcion", "Codigo",
           "Fragmento_parafraseado", "Motivo", "Cita_anterior",
           "Cand1_linea", "Cand1_parecido", "Cand1_texto",
           "Cand2_linea", "Cand2_parecido", "Cand2_texto",
           "Cand3_linea", "Cand3_parecido", "Cand3_texto",
           "ALERTA_otra_entrevista",
           "Revisor", "DECISION", "CITA_ELEGIDA", "LINEA_ELEGIDA", "OBSERVACION"]
    ws.append(cab)
    for i, (nombre, fila, propia, alerta) in enumerate(registros, start=1):
        motivo = "SIN_CITA" if fila["ESTADO"].strip() != "VER" else "CITA_NO_LITERAL"
        r = [i, nombre, int(fila["linea_csv"]), fila["ID_evidencia"], corto(fila["transcripcion"]),
             fila["Codigo"], fila["Fragmento_parafraseado"], motivo,
             fila["CITA_LITERAL"] if motivo == "CITA_NO_LITERAL" else ""]
        for j in range(3):
            if j < len(propia):
                r += [propia[j][1], propia[j][0], propia[j][2]]
            else:
                r += ["", "", ""]
        r += [alerta, "", "", "", "", ""]
        ws.append(r)

    fuente = "Arial"
    borde_cab = PatternFill("solid", start_color="1F3864")
    amarillo = PatternFill("solid", start_color="FFF2CC")
    alerta_fill = PatternFill("solid", start_color="F8CBAD")
    anchos = {"N": 5, "Archivo_CSV": 16, "Fila_CSV": 8, "ID_evidencia": 9, "Transcripcion": 11,
              "Codigo": 24, "Fragmento_parafraseado": 46, "Motivo": 16, "Cita_anterior": 34,
              "Cand1_linea": 8, "Cand1_parecido": 9, "Cand1_texto": 58,
              "Cand2_linea": 8, "Cand2_parecido": 9, "Cand2_texto": 58,
              "Cand3_linea": 8, "Cand3_parecido": 9, "Cand3_texto": 58,
              "ALERTA_otra_entrevista": 44, "Revisor": 13, "DECISION": 15,
              "CITA_ELEGIDA": 50, "LINEA_ELEGIDA": 9, "OBSERVACION": 34}
    for c, nombre in enumerate(cab, start=1):
        celda = ws.cell(row=1, column=c)
        celda.font = Font(name=fuente, bold=True, color="FFFFFF", size=10)
        celda.fill = borde_cab
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = anchos[nombre]
    entrada = {cab.index(n) + 1 for n in ("Revisor", "DECISION", "CITA_ELEGIDA", "LINEA_ELEGIDA", "OBSERVACION")}
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.font = Font(name=fuente, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.column in entrada:
                c.fill = amarillo
        if fila[cab.index("ALERTA_otra_entrevista")].value:
            fila[cab.index("ALERTA_otra_entrevista")].fill = alerta_fill
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    col = get_column_letter(cab.index("DECISION") + 1)
    dv.add(f"{col}2:{col}{len(registros) + 1}")
    ws.freeze_panes = "H2"
    ws.auto_filter.ref = ws.dimensions

    ins = wb.create_sheet("Instrucciones")
    texto = [
        ("Cómo revisar (una fila = un fragmento sin cita literal válida)", True),
        ("1. Lee Fragmento_parafraseado y Codigo. Mira Cita_anterior (versión depurada del texto) para orientarte.", False),
        ("2. Lee las 3 candidatas: son fragmentos EXACTOS de lo que dijo el entrevistado (nunca del entrevistador).", False),
        ("3. Si una candidata RESPALDA el código (no basta que hable del mismo tema), escribe su DECISION: CAND1, CAND2 o CAND3.", False),
        ("4. Si ninguna sirve pero tú encuentras la frase en la transcripción: DECISION = LITERAL y pega EXACTA la frase en CITA_ELEGIDA.", False),
        ("5. Si no existe respaldo en esa entrevista: DECISION = NO_LOCALIZADA y explica en OBSERVACION (ver ALERTA_otra_entrevista).", False),
        ("6. Escribe tu apellido en Revisor: es la constancia de quién revisó; no firmes filas que no leíste.", False),
        ("7. Celdas amarillas = las únicas que se editan. Columna naranja = otra entrevista respalda mejor (útil para C2).", False),
        ("8. Al terminar: reanclar_citas_C1.py aplicar <este archivo> --simular, luego sin --simular, luego verificar_citas_C1.py.", False),
        ("", False),
        ("Ejemplo de filas ya revisadas (solo ilustrativo, no se aplica):", True),
    ]
    for t, negrita in texto:
        ins.append([t])
        ins.cell(row=ins.max_row, column=1).font = Font(name=fuente, size=10, bold=negrita)
    ins.append(["Revisor", "DECISION", "CITA_ELEGIDA", "OBSERVACION"])
    ins.append(["Apellido", "CAND2", "", "La candidata 2 respalda el código (habla de la frecuencia semanal)."])
    ins.append(["Apellido", "LITERAL", "frase exacta copiada de la transcripción", "Ninguna candidata servía; la frase está más adelante."])
    ins.append(["Apellido", "NO_LOCALIZADA", "", "No hay respaldo aquí; ENTR-08 línea 41 lo dice mejor (posible C2)."])
    for fila in ins.iter_rows(min_row=ins.max_row - 3):
        for c in fila:
            c.font = Font(name=fuente, size=10, bold=(c.row == ins.max_row - 3))
    ins.column_dimensions["A"].width = 120
    for letra in "BCD":
        ins.column_dimensions[letra].width = 40

    HOJA_SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(HOJA_SALIDA)
    n_alerta = sum(1 for r in registros if r[3])
    n_sin = sum(1 for r in registros if r[1]["ESTADO"].strip() != "VER")
    print(f"{len(registros)} fragmentos pendientes ({n_sin} sin cita, {len(registros) - n_sin} con cita no literal);"
          f" {n_alerta} con alerta de otra entrevista")
    print(f"-> {HOJA_SALIDA.relative_to(RAIZ)}")


# --------------------------------------------------------------- aplicar
def leer_hoja(ruta):
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(ruta, data_only=True)["Revision"]
        cab = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        return [{cab[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(fila)}
                for fila in ws.iter_rows(min_row=2, values_only=True) if fila[0] is not None]
    crudo = ruta.read_bytes().decode("utf-8-sig", errors="replace")
    delim = max([";", ",", "\t"], key=lambda d: crudo.split("\n", 1)[0].count(d))
    return list(csv.DictReader(crudo.splitlines(), delimiter=delim))


def localizar(cita, turnos, pista):
    if len(cita) < 15:
        return None
    if str(pista).strip().isdigit():
        for n, l, o in turnos:
            if n == int(pista) and cita in l[o:]:
                return n
    for n, l, o in turnos:
        if cita in l[o:]:
            return n
    return None


def aplicar(ruta_hoja, simular):
    hoja = leer_hoja(ruta_hoja)
    por_clave = {(r["Archivo_CSV"], str(int(float(r["Fila_CSV"])))): r for r in hoja}
    turnos_por_archivo = cargar_turnos()
    resumen = {"aplicadas (VER)": 0, "NO_LOCALIZADA": 0, "sin decisión (no se tocan)": 0}
    errores = []
    for nombre, ruta in ARCHIVOS_CSV.items():
        filas, campos, bom, crlf = leer_csv(ruta)
        for fila in filas:
            rev = por_clave.get((nombre, str(fila["linea_csv"]).strip()))
            if not rev or not rev.get("DECISION", "").strip():
                if rev:
                    resumen["sin decisión (no se tocan)"] += 1
                continue
            dec, revisor = rev["DECISION"].strip().upper(), rev.get("Revisor", "").strip()
            etiqueta = f"{nombre} fila {fila['linea_csv']} ({fila['ID_evidencia']} {fila['Codigo']})"
            if not revisor:
                errores.append(f"{etiqueta}: falta el nombre del Revisor -> no se aplica")
                continue
            if dec == "NO_LOCALIZADA":
                fila.update({"CITA_LITERAL": "", "LINEA_TRANSCRIPCION": "", "ESTADO": "NO_LOCALIZADA",
                             "REVISOR_CITA": revisor})
                resumen["NO_LOCALIZADA"] += 1
                continue
            if dec in ("CAND1", "CAND2", "CAND3"):
                cita, pista = rev.get(f"Cand{dec[-1]}_texto", ""), rev.get(f"Cand{dec[-1]}_linea", "")
            elif dec == "LITERAL":
                cita, pista = rev.get("CITA_ELEGIDA", ""), rev.get("LINEA_ELEGIDA", "")
            else:
                errores.append(f"{etiqueta}: DECISION '{dec}' no reconocida -> no se aplica")
                continue
            n = localizar(cita, turnos_por_archivo[fila["transcripcion"]], pista)
            if n is None:
                errores.append(f"{etiqueta}: la cita NO es una subcadena exacta (>=15 car.) de un turno del "
                               f"entrevistado en {corto(fila['transcripcion'])} -> no se aplica")
                continue
            fila.update({"CITA_LITERAL": cita, "LINEA_TRANSCRIPCION": str(n), "ESTADO": "VER",
                         "REVISOR_CITA": revisor})
            resumen["aplicadas (VER)"] += 1
        if not simular:
            escribir_csv(ruta, filas, campos, bom, crlf)

    print("Resultado" + (" (SIMULACIÓN: no se escribió nada)" if simular else ""))
    for k, v in resumen.items():
        print(f"  {k}: {v}")
    if errores:
        print(f"\n{len(errores)} filas rechazadas:")
        for e in errores:
            print("  -", e)
    print("\nSiguiente paso: python3 07_Datos/scripts/plan_mejora/verificar_citas_C1.py")
    return 1 if errores else 0


def main():
    if len(sys.argv) < 2 or sys.argv[1] not in ("proponer", "aplicar"):
        sys.exit(__doc__)
    if sys.argv[1] == "proponer":
        proponer()
        return 0
    args = [a for a in sys.argv[2:] if not a.startswith("--")]
    if not args:
        sys.exit("Indica la hoja: reanclar_citas_C1.py aplicar HOJA.xlsx [--simular]")
    return aplicar(args[0], "--simular" in sys.argv)


if __name__ == "__main__":
    sys.exit(main())
