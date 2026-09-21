#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
matriz_cobertura_B4.py  ·  B4 · SIMPA_ISR401

Criterio de B4: "Matriz 13 x 16 con cita por celda" y "limitaciones declaradas".

Las 13 preguntas son el banco de la guía A.2.1 (09_Etica/A02_Instrumentos_Recoleccion.pdf).
Para cada una de las 16 entrevistas, la celda dice si la pregunta se cubrió y trae la
cita LITERAL del entrevistado que lo prueba:

    P        = PREGUNTADA: el entrevistador la formuló y la persona respondió.
    E        = ESPONTÁNEA: no se preguntó, pero la persona lo abordó.
    PARCIAL  = responde solo una parte de la pregunta (se indica cuál en la nota).
    -        = NO ABORDADA: ni preguntada ni mencionada (sin cita).

QUÉ ES HUMANO Y QUÉ ES AUTOMÁTICO
    · MAPA (abajo) es la decisión de qué respuesta cubre qué pregunta y con qué estado.
      Fue hecha leyendo las 16 transcripciones (asistida por Claude; declarar en B2) y
      debe poder revisarse celda por celda en matriz_cobertura_B4.csv.
    · Este script busca cada cita indicada en la MAPA dentro de la respuesta señalada, la copia
      LITERAL y comprueba que sea subcadena exacta de un turno del entrevistado (no del
      entrevistador). Si una cita no se comprueba, el script falla y no escribe nada.

USO (desde la raíz del repositorio)
    python3 07_Datos/scripts/plan_mejora/matriz_cobertura_B4.py
Salidas
    07_Datos/datos_procesados/matriz_cobertura_B4.csv      (208 filas: pregunta x entrevista)
    07_Datos/resultados/matriz_cobertura_B4.xlsx            (matriz 13x16, citas, limitaciones)
    07_Datos/resultados/matriz_cobertura_B4.md              (resumen y limitaciones)
"""
import csv
import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import citas_lib as L  # noqa: E402

RAIZ = L.RAIZ
SALIDA_CSV = RAIZ / "07_Datos" / "datos_procesados" / "matriz_cobertura_B4.csv"
SALIDA_XLSX = RAIZ / "07_Datos" / "resultados" / "matriz_cobertura_B4.xlsx"
SALIDA_MD = RAIZ / "07_Datos" / "resultados" / "matriz_cobertura_B4.md"

PREGUNTAS = {
    1: "¿Con qué frecuencia revisa el estado general del cultivo y qué aspectos observa en esas revisiones?",
    2: "¿Cómo planifica las labores por parcela o lote y cómo se distribuye el trabajo entre los equipos?",
    3: "¿Cómo planifica la aplicación de productos fitosanitarios y cómo controla los periodos de reingreso y de carencia?",
    4: "¿Cómo registra la ocurrencia de plagas y enfermedades y su evolución en el tiempo?",
    5: "¿Qué señales observa para determinar si una labor se ejecutó correctamente?",
    6: "¿Cómo verifica el cumplimiento del trabajo del personal de campo y con qué criterios lo evalúa?",
    7: "¿Cómo se maneja el registro y reporte hacia los organismos de control del sector y qué información exigen?",
    8: "¿Cómo se coordina la cosecha y el envío de la fruta a la extractora?",
    9: "¿Cómo obtiene y registra la información climática y cómo influye en sus decisiones?",
    10: "¿Qué registros mantiene actualmente en papel o en hojas de cálculo?",
    11: "¿Qué información querría consultar en campo desde un dispositivo móvil?",
    12: "¿Qué tan útil considera un apoyo automático que, a partir de una fotografía, sugiera el problema que presenta la planta?",
    13: "¿Cuál es el principal problema que enfrenta hoy en la gestión del cultivo o del personal?",
}
ENTREVISTAS = [f"ENTR-{i:02d}" for i in range(1, 17)]

# MAPA[entrevista] = [(pregunta, estado, línea de la respuesta[, inicio_cita, fin_cita, nota])]
# La línea es la del PRIMER párrafo de la respuesta del entrevistado en el archivo de transcripción.
MAPA = {
    "ENTR-01": [
        (1, "P", 11, "Entonces la función del administrador es recorrer todos los trabajos", "cumpliendo con su funcion.", ""),
        (2, "E", 11, "se tiene un equipo de polinizacion", "un equipo de labores.", ""),
        (5, "E", 75, "Si si tienes una fruta eh en buena en buena calidad", "no está cumpliendo con su trabajo", "Respuesta a la pregunta por el desempeño de los trabajadores."),
        (6, "P", 75, "Eh por rendimiento y producción,", "flores que han polinizado en el día", ""),
        (9, "P", 59, "siempre medimos el tema de lluvias", "a nivel general.", ""),
        (11, "P", 67, "Eh el tema uno de plagas, cómo reconocer la plaga", "las etapas de de la plaga", "Pregunta sobre información deseada del análisis de imágenes con IA."),
        (12, "P", 63, "Sí, en el tema de de de plagas y enfermedades sería súper interesante", "la familia de la plaga", "Pregunta INDUCIDA: el entrevistador propone el análisis de imágenes con IA (línea 61)."),
        (13, "P", 83, "En en el tema de de a nivel general el problema es la mano de obra.", "", ""),
    ],
    "ENTR-02": [
        (1, "P", 15, "revisamos diariamente todas las actividades", "como la chapia y la corona.", ""),
        (2, "E", 43, "en esta plantación que es pequeña tenemos dos", "dan la vuelta todo alrededor", "Sobre sanidad; no describe la planificación por lote."),
        (3, "PARCIAL", 27, "Eso tenemos que hacerlo por lo menos repetidamente cada semana", "hasta que la plaga termine.", "Solo la frecuencia de aplicación; nada sobre reingreso ni carencia."),
        (6, "P", 95, "Y ahí nos damos cuenta si el trabajador", "no está haciendo nada.", ""),
        (9, "P", 71, "Para nosotros poder ver cuánto milímetro de de lluvia ha caído de repente.", "", ""),
        (10, "PARCIAL", 71, "Ahí llevamos un registro, ¿no?", "", "Registro de lluvia; no dice si es en papel o en hoja de cálculo."),
        (13, "P", 19, "Sí, los problemas más comunes en este caso que son... prácticamente la nutrición.", "", ""),
    ],
    "ENTR-03": [
        (5, "P", 33, "La mayor, es visible, por lo que es un talco industrial más una hormona", "primer momento de aplicación.", ""),
        (6, "P", 45, "Con un rastreador GPS.", "", "Respuesta a '¿mediante qué control?'."),
        (9, "P", 37, "El clima influye porque la palma necesita horas luz", "su maduración.", ""),
        (11, "P", 49, "Que el rastreador GPS tenga marcaciones", "polinizadas diariamente.", ""),
        (13, "P", 29, "Los problemas que más se frecuentan es el problema laboral.", "", ""),
    ],
    "ENTR-04": [
        (7, "E", 63, "Actualmente muchos de los palmeros en sí no manejan esta guía de remisión", "exigido por Agrocalidad.", "La guía de remisión la exige Agrocalidad; no se preguntó por reportes a organismos de control."),
        (8, "P", 39, "Entonces, nosotros normalmente recomendamos que sea en las 24 horas", "hasta 48, 30 y un poquito más.", ""),
        (10, "PARCIAL", 59, "nosotros tenemos un sistema interno en la empresa", "en un ticket.", "Registro en sistema interno de la extractora, no en papel u hoja de cálculo."),
        (11, "P", 79, "que se pueda llevar quizás cuando suban el conteo de racimos", "en este caso esta de aquí también.", "Respuesta sobre una aplicación en el celular de la hacienda."),
        (13, "P", 31, "Ya, el principal problema que nosotros tenemos actualmente es el tema de fruta verde.", "", ""),
    ],
    "ENTR-05": [
        (2, "P", 15, "El administrador.", "", "Respuesta corta a '¿quién le dice en qué lote o parte le toca trabajar?'."),
        (4, "PARCIAL", 53, "Cuando nosotros la vemos, inmediatamente le decimos a él", "ingeniería y todo.", "Cómo se avisa de una planta enferma; no describe un registro ni su evolución."),
        (6, "PARCIAL", 23, "eso se lleva un conteo de las flores", "se le da el conteo al encargado", "Visión del trabajador: entrega el conteo; no dice cómo se verifica."),
        (10, "P", 27, "Claro, nosotros le decimos, lo van a tener en una libreta", "que el hombre nos pague.", ""),
        (11, "P", 77, "Claro, porque ya sé lo que estoy ganándome", "cuánto me sacaría en la quincena.", ""),
        (13, "P", 49, "Pues ahorita por el modo de la cosecha, que ahorita está muy difícil", "por lo que es primera cosecha", ""),
    ],
    "ENTR-06": [
        (2, "P", 27, "Lote uno y lote dos.", "", "El jefe asigna el lote (respuesta 'El jefe.', línea 33, demasiado corta para citar)."),
        (10, "P", 43, "Sí, la plantación que hizo. Lo hace de un cuadernito.", "", ""),
        (11, "P", 67, "Por ejemplo, que le cuente las plantas.", "Contar las plantas, todo eso.", ""),
    ],
    "ENTR-07": [
        (1, "E", 47, "como las visitas yo las hago dos veces por semana", "para poder verificar.", "Frecuencia de las visitas de verificación."),
        (2, "E", 51, "Y para saber qué lote es, eso también me lo proporciona el jefe de campo", "para revisar.", ""),
        (4, "P", 43, "lo que se hace es tomar una evidencia fotográfica", "qué tipo de enfermedad tiene", "Registro por foto de la planta enferma; no describe evolución."),
        (5, "PARCIAL", 35, "todo eso se corrige al momento de revisar.", "", "Errores en el conteo que se corrigen al revisar."),
        (6, "E", 19, "llego a la finca, localizo al jefe de campo", "en la finca.", ""),
        (10, "P", 27, "la libreta realmente la lleva al jefe de campo", "administradores.", ""),
        (11, "P", 71, "Sería interesante que se pudiese guardar los seguimientos con evidencia fotográfica", "trabajos.", ""),
        (13, "P", 39, "diría que los seguimientos, que no suelen llegar los reportes lo suficientemente especificados", "jefe de campo.", ""),
    ],
    "ENTR-08": [
        (5, "PARCIAL", 79, "primero observar bien la la parte basal de la planta para ver si hay desprendimiento.", "", "Señal de racimo listo para cortar; no de labor ejecutada."),
        (6, "E", 51, "Entonces nosotros recurrimos a la visita técnica para ver cuál es el problema", "el grupo de cosecha o el cultivo.", ""),
        (11, "P", 99, "En sí nosotros recomendamos ver este... eh sí que mande bien", "hacemos monitoreo de insectos.", "Qué mostraría la aplicación de fotos."),
        (12, "P", 95, "En sí sí nos serviría porque", "durante el cultivo.", ""),
        (13, "P", 71, "Lo más básico siempre va a ser la mala selección de frutos", "el único error que cometen.", ""),
    ],
    "ENTR-09": [
        (2, "P", 19, "quien decide qué hace cada persona normalmente va de la mano del administrador de la hacienda", "repartir las labores.", ""),
        (6, "P", 23, "Bueno, normalmente se lleva mediante un conteo.", "", ""),
        (10, "PARCIAL", 31, "Normalmente se usa una aplicación", "se controla el recorrido de cada persona.", "Habla de una aplicación (Avenza Maps), no de papel u hojas de cálculo."),
        (11, "P", 47, "En este caso llevaría un control de recorrido por parte del personal", "labores que se van haciendo", ""),
        (12, "P", 43, "Sería una buena ayuda porque hay muchos técnicos.", "", ""),
        (13, "P", 27, "Normalmente, hoy en día, un problema muy grande es la inseguridad.", "", ""),
    ],
    "ENTR-10": [
        (6, "P", 67, "O sea, que la aplicación, cuando termines el recorrido, te muestra el recorrido", "todo eso.", "Pregunta sobre registrar por GPS el recorrido del personal."),
        (10, "P", 31, "Aplicaciones como WhatsApp.", "", "Respuesta a '¿papel, memoria o una aplicación?'."),
        (11, "P", 51, "Por ejemplo, el pronóstico, la semana de polinización, puede ser.", "", ""),
        (12, "PARCIAL", 59, "Más que todo, la información y fotos.", "", "Ante un aviso de posible plaga; no valora un diagnóstico automático."),
    ],
    "ENTR-11": [
        (6, "P", 87, "La verdad es que me parece muy bien.", "", "Pregunta sobre registrar por GPS el recorrido del personal."),
        (10, "P", 39, "Con una estala, se llevan los datos; día a día se va a organizar.", "", "Luego confirma que es 'Mediante papel.' (línea 43)."),
        (11, "P", 67, "La información, primero la foto y el informe de los correos.", "", ""),
        (12, "PARCIAL", 71, "Primero, una imagen donde me compruebe sobre ese problema", "resolver el problema.", "Ante un aviso de posible plaga; pide imagen y comprobar en campo."),
    ],
    "ENTR-12": [
        (2, "P", 23, "primero yo considero que hay que hacer una inspección en la planta", "del día completo.", ""),
        (10, "PARCIAL", 39, "Como muy efectivas. Nos ayudan muchísimo", "se podría decir.", "Habla de aplicaciones de control en general."),
        (11, "P", 63, "Por ejemplo, el estado de las plantaciones.", "", ""),
        (13, "P", 31, "Bueno, consideraría las enfermedades que frecuentemente enfrenta el cultivo.", "", ""),
    ],
    "ENTR-13": [
        (2, "P", 21, "en una finca pequeña, de poco personal, pues siempre a veces decide el dueño o un encargado.", "", ""),
        (6, "P", 25, "en las fincas pequeñas, muchos llevan registros empíricos.", "", ""),
        (8, "P", 29, "Nosotros específicamente tenemos los técnicos de campo", "periodos de cosecha que ellos tienen.", "Es comprador: programa la cosecha con los proveedores."),
        (9, "P", 49, "el agricultor siempre se fija más en ver cómo está el ambiente", "programar alguna actividad que se tenga en campo.", ""),
        (10, "P", 55, "en fincas pequeñas a lo mucho llevan una agenda, una libreta.", "", ""),
        (11, "P", 77, "toda la información de las labores que están pendientes de hacer", "las labores que ya se han hecho.", ""),
        (12, "P", 73, "La tecnología debe estar ligada actualmente a la agricultura.", "", ""),
        (13, "P", 39, "primeramente, a veces cuando la gente es un poquito incumplida.", "", ""),
    ],
    "ENTR-14": [
        (2, "P", 15, "hay dos tipos de organizaciones.", "qué es lo que va a realizar en el día de trabajo.", ""),
        (6, "P", 19, "el control se lleva mediante una bitácora.", "esa semana.", ""),
        (10, "E", 47, "esto se hace mediante una bitácora semanal", "la siguiente semana", "Bitácora semanal (documento firmado); no indica si es en papel o en hoja de cálculo."),
        (11, "P", 47, "Especialmente las labores que vamos a hacer a la semana", "", ""),
        (12, "P", 43, "Parece interesante.", "Muy interesante me parece.", ""),
        (13, "P", 27, "El problema más común son los productos.", "aplicado en su totalidad.", ""),
    ],
    "ENTR-15": [
        (2, "PARCIAL", 26, "recurriría principalmente a la persona encargada del campo", "administrador", "A quién consulta lo que se hizo en un lote; no cómo se planifica."),
        (6, "P", 42, "creo que no tanto se basaría en el seguimiento", "más se debería enfocar como en los trabajos que se realizó.", "Pregunta sobre registrar por GPS el recorrido del personal."),
        (10, "P", 14, "Por lo general solo hacen con papel.", "", ""),
        (11, "P", 30, "Sería lo lo primordial, atender lo que está atrasado", "en una buena condición.", ""),
        (12, "PARCIAL", 34, "al recibir la notificación iría directamente al sector donde está pasando aquello.", "", "Ante un aviso de posible plaga; no valora un diagnóstico automático."),
        (13, "PARCIAL", 18, "Sí, cuando se trata de sacar las herramientas o el equipo", "te atrasa en la labor", "Problema por información que llega tarde; no es 'el principal problema' de la gestión."),
    ],
    "ENTR-16": [
        (2, "P", 18, "dependerá mucho de la organización de una finca, si es empresarial o si es familiar.", "", ""),
        (6, "P", 22, "Dependerá mucho de los dos ámbitos", "si es empresarial o si es doméstico.", ""),
        (10, "PARCIAL", 30, "Una optimización de procesos en el área agrícola vendría muy bien ahora en la actualidad.", "", "Habla de soluciones informáticas en general."),
        (11, "P", 42, "lo que más interesaría en ese entonces en su implementación", "cosechas, cultivos y manejos del fruto.", ""),
        (13, "P", 26, "el problema más común que existe son las enfermedades.", "", ""),
    ],
}
NOTA_NO = {
    3: "No se formuló; solo se pregunta por reingreso o carencia en la guía A.2.",
    7: "No se formuló ni se abordó el reporte a organismos de control.",
}
LIMITACIONES = [
    "La guía A.2 (banco de 13 preguntas) se creó el 30/07/2026, después de las rondas 1 y 2 (ENTR-01 a 03 del 23/05 y ENTR-04 a 08 del 28/07). Esas ocho entrevistas no siguieron el banco: usaron preguntas por rol (trabajador, técnico, administrador). Las de 31/08 y 01/09 (ENTR-09 a 16) usaron una guía corta distinta.",
    "Las preguntas 3 (planificación de fitosanitarios, reingreso y carencia) y 7 (reporte a organismos de control) no se formularon en ninguna entrevista; solo hay menciones incidentales (ENTR-02 sobre la frecuencia de aplicación, ENTR-04 sobre la guía de remisión que exige Agrocalidad).",
    "La pregunta 12 (utilidad de un apoyo automático por fotografía) se formuló de forma directa solo en algunas entrevistas; en ENTR-01 la introduce el entrevistador (pregunta inducida) y en varias solo se preguntó por un aviso de plaga en una aplicación (estado PARCIAL).",
    "'Cobertura' significa que la persona habló del tema con una cita literal; no mide la calidad de la respuesta ni la profundidad. Las celdas PARCIAL indican qué parte de la pregunta queda sin responder.",
    "Las entrevistas 09 a 16 (ronda 3) son a profesionales y estudiantes de carreras afines, no a personal operativo; la comparación entre rondas no es entre perfiles equivalentes.",
    "En ENTR-02 se excluyó el tramo posterior a [16:04], que no es fiable hasta compararlo con el audio (ver hallazgo de B1); las preguntas que solo se responden ahí no cuentan como cubiertas.",
    "La asignación de cada celda la hizo un solo analista (asistido por Claude, declarado en B2) y no fue doblemente codificada; es una decisión revisable en matriz_cobertura_B4.csv.",
    "ENTR-05 y ENTR-06 estaban intercambiadas en las transcripciones originales; las citas de esta matriz siguen las retranscripciones vigentes.",
]


def construir():
    turnos = {}
    for e in ENTREVISTAS:
        f = next(L.TRANSCRIPCIONES.glob(f"*_{e}_Transcripcion.md"))
        turnos[e] = (f.name, L.turnos_de_archivo(f))
    filas = []
    for e in ENTREVISTAS:
        nombre, tr = turnos[e]
        por_q = {m[0]: m for m in MAPA.get(e, [])}
        for q in range(1, 14):
            if q not in por_q:
                filas.append({"pregunta": q, "entrevista": e, "estado": "NO_ABORDADA", "cita_literal": "",
                              "linea": "", "transcripcion": nombre, "nota": NOTA_NO.get(q, "")})
                continue
            _, est, linea, ini, fin, nota = por_q[q]
            estado = {"P": "PREGUNTADA", "E": "ESPONTANEA", "PARCIAL": "PARCIAL"}[est]
            hit = None
            for n, lin, off in tr:
                txt = lin[off:]
                a = txt.find(ini)
                if a < 0 or not (linea <= n <= linea + 8):
                    continue
                if fin:
                    b = txt.find(fin, a)
                    if b < 0:
                        continue
                    b += len(fin)
                else:
                    b = a + len(ini)
                hit = (n, txt[a:b].strip())
                break
            if not hit:
                raise SystemExit(f"NO se encontró la cita de {e} pregunta {q} (desde línea {linea}): {ini[:50]}")
            filas.append({"pregunta": q, "entrevista": e, "estado": estado, "cita_literal": hit[1],
                          "linea": hit[0], "transcripcion": nombre, "nota": nota})
    return filas


def verificar(filas):
    """13 x 16 completo, y toda cita es subcadena EXACTA de un turno del entrevistado."""
    errores = []
    if len(filas) != 13 * 16 or {(f["pregunta"], f["entrevista"]) for f in filas} != {(q, e) for q in range(1, 14) for e in ENTREVISTAS}:
        errores.append("La matriz no es 13 x 16 completa.")
    cache = {}
    for f in filas:
        if f["estado"] == "NO_ABORDADA":
            if f["cita_literal"]:
                errores.append(f"{f['entrevista']} P{f['pregunta']}: NO_ABORDADA con cita")
            continue
        tr = cache.setdefault(f["transcripcion"], L.turnos_de_archivo(L.TRANSCRIPCIONES / f["transcripcion"]))
        if not f["cita_literal"] or not any(n == f["linea"] and f["cita_literal"] in l[o:] for n, l, o in tr):
            errores.append(f"{f['entrevista']} P{f['pregunta']}: la cita no es literal en la línea {f['linea']}")
    return errores


def escribir(filas):
    campos = ["pregunta", "entrevista", "estado", "cita_literal", "linea", "transcripcion", "nota"]
    SALIDA_CSV.parent.mkdir(parents=True, exist_ok=True)
    with open(SALIDA_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=campos, delimiter=";")
        w.writeheader()
        w.writerows(filas)

    cont = Counter(f["estado"] for f in filas)
    por_q = {q: Counter(f["estado"] for f in filas if f["pregunta"] == q) for q in range(1, 14)}
    md = ["# Matriz de cobertura: pregunta x entrevista (B4)", "",
          "Banco de 13 preguntas de la guía A.2.1 x 16 entrevistas. Cada celda cubierta trae una cita literal "
          "(ver `07_Datos/datos_procesados/matriz_cobertura_B4.csv` y `matriz_cobertura_B4.xlsx`).", "",
          f"Celdas: 208 · PREGUNTADA {cont['PREGUNTADA']} · ESPONTANEA {cont['ESPONTANEA']} · "
          f"PARCIAL {cont['PARCIAL']} · NO_ABORDADA {cont['NO_ABORDADA']}", "",
          "| Pregunta | Preguntada | Espontánea | Parcial | No abordada |", "|---|---|---|---|---|"]
    for q in range(1, 14):
        c = por_q[q]
        md.append(f"| {q}. {PREGUNTAS[q][:70]}… | {c['PREGUNTADA']} | {c['ESPONTANEA']} | {c['PARCIAL']} | {c['NO_ABORDADA']} |")
    md += ["", "## Limitaciones declaradas", ""] + [f"{i}. {t}" for i, t in enumerate(LIMITACIONES, 1)]
    SALIDA_MD.write_text("\n".join(md) + "\n", encoding="utf-8")

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    fuente = "Arial"
    color = {"PREGUNTADA": "C6E0B4", "ESPONTANEA": "DDEBF7", "PARCIAL": "FFE699", "NO_ABORDADA": "F2F2F2"}
    sigla = {"PREGUNTADA": "P", "ESPONTANEA": "E", "PARCIAL": "PARCIAL", "NO_ABORDADA": "—"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz 13x16"
    ws.append(["Pregunta de la guía A.2.1"] + [e.replace("ENTR-", "E") for e in ENTREVISTAS])
    por_celda = {(f["pregunta"], f["entrevista"]): f for f in filas}
    for q in range(1, 14):
        fila = [f"{q}. {PREGUNTAS[q]}"]
        for e in ENTREVISTAS:
            f = por_celda[(q, e)]
            fila.append(f"[{sigla[f['estado']]}] «{f['cita_literal']}»" if f["cita_literal"] else "—")
        ws.append(fila)
    ws.append([])
    ws.append(["Leyenda: [P] preguntada · [E] espontánea · [PARCIAL] responde solo una parte · — no abordada. "
               "Cada cita es literal (hoja 'Citas' indica línea y transcripción)."])
    for c in range(1, 18):
        h = ws.cell(row=1, column=c)
        h.font = Font(name=fuente, bold=True, color="FFFFFF", size=10)
        h.fill = PatternFill("solid", start_color="1F3864")
        h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = 46 if c == 1 else 34
    for r in range(2, 15):
        ws.row_dimensions[r].height = 150
        ws.cell(row=r, column=1).font = Font(name=fuente, bold=True, size=9)
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        for c, e in enumerate(ENTREVISTAS, start=2):
            cel = ws.cell(row=r, column=c)
            cel.font = Font(name=fuente, size=8)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.fill = PatternFill("solid", start_color=color[por_celda[(r - 1, e)]["estado"]])
    ws.cell(row=16, column=1).font = Font(name=fuente, italic=True, size=9)
    ws.freeze_panes = "B2"

    ws2 = wb.create_sheet("Citas")
    ws2.append(["pregunta", "entrevista", "estado", "cita_literal", "linea", "transcripcion", "nota"])
    for f in filas:
        ws2.append([f[k] for k in ["pregunta", "entrevista", "estado", "cita_literal", "linea", "transcripcion", "nota"]])
    for c, w in enumerate([9, 11, 14, 90, 7, 34, 60], start=1):
        ws2.column_dimensions[get_column_letter(c)].width = w
        ws2.cell(row=1, column=c).font = Font(name=fuente, bold=True)
    for fila in ws2.iter_rows(min_row=2):
        for c in fila:
            c.font = Font(name=fuente, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.auto_filter.ref = ws2.dimensions
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Resumen y limitaciones")
    ws3.append([f"Celdas: 208 · PREGUNTADA {cont['PREGUNTADA']} · ESPONTANEA {cont['ESPONTANEA']} · "
                f"PARCIAL {cont['PARCIAL']} · NO_ABORDADA {cont['NO_ABORDADA']}"])
    ws3.append([])
    ws3.append(["Pregunta", "Preguntada", "Espontánea", "Parcial", "No abordada"])
    for q in range(1, 14):
        c = por_q[q]
        ws3.append([f"{q}. {PREGUNTAS[q]}", c["PREGUNTADA"], c["ESPONTANEA"], c["PARCIAL"], c["NO_ABORDADA"]])
    ws3.append([])
    ws3.append(["Limitaciones declaradas"])
    for i, t in enumerate(LIMITACIONES, 1):
        ws3.append([f"{i}. {t}"])
    ws3.column_dimensions["A"].width = 130
    for r in ws3.iter_rows():
        for c in r:
            c.font = Font(name=fuente, size=10, bold=(c.row in (3,) or c.value == "Limitaciones declaradas"))
            c.alignment = Alignment(wrap_text=True, vertical="top")
    SALIDA_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA_XLSX)
    return cont, por_q


def main():
    filas = construir()
    errores = verificar(filas)
    if errores:
        print("ERRORES:\n  " + "\n  ".join(errores))
        return 1
    cont, por_q = escribir(filas)
    print(f"Matriz 13 x 16 completa: {dict(cont)}")
    print("Preguntadas (P) por pregunta:", {q: por_q[q]["PREGUNTADA"] for q in range(1, 14)})
    print("Cubiertas (P+E+PARCIAL) por pregunta:", {q: 16 - por_q[q]["NO_ABORDADA"] for q in range(1, 14)})
    print("VERIFICACIÓN: todas las citas son literales. OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
